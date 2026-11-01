"""Synthetic fixtures, built in code: no family's file is ever one of them.

The tests directory is on sys.path (pytest's prepend import mode), so tests import `fixtures`.
"""

from __future__ import annotations

import io
import shutil
import subprocess
from typing import Any

import pytest
from PIL import Image

POD = "pod_0aaaaaaaaaaaaaaaaaaaaaaaaa"
OTHER_POD = "pod_0bbbbbbbbbbbbbbbbbbbbbbbbb"
ATTACHMENT = "atch_0ccccccccccccccccccccccccc"
BUCKET = "jadwigo-uploads-test"

#: An A4 page in points, which pdftoppm renders at 150 dpi as 1240 x 1755 px.
A4 = (595, 842)


def event(kind: str, media_type: str, **overrides: str) -> dict[str, str]:
    base = {
        "bucket": BUCKET,
        "key": f"pods/{POD}/attachments/{ATTACHMENT}.{kind}",
        "media_type": media_type,
        "pod": POD,
        "attachment_id": ATTACHMENT,
    }
    return {**base, **overrides}


def keys(s3: Any) -> list[str]:
    """Every key in the bucket, sorted."""
    return sorted(o["Key"] for o in s3.list_objects_v2(Bucket=BUCKET).get("Contents", []))


def pdf_bytes(pages: int, *, user_password: str | None = None, owner_password: str | None = None) -> bytes:
    """A PDF of blank A4 pages, encrypted when a password is given."""
    from pypdf import PdfWriter

    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=A4[0], height=A4[1])
    if user_password is not None or owner_password is not None:
        writer.encrypt(
            user_password=user_password or "",
            owner_password=owner_password,
            algorithm="RC4-128",
        )
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def docx_bytes(pages: int = 2) -> bytes:
    """A worksheet-like DOCX, one page break between each of its pages."""
    from docx import Document

    doc = Document()
    for n in range(1, pages + 1):
        doc.add_heading(f"Zadanie {n}", level=1)
        doc.add_paragraph("Ile to jest 7 + 5? Zapisz wynik w kratce.")
        if n < pages:
            doc.add_page_break()
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def exif_bytes(orientation: int = 1) -> bytes:
    exif = Image.Exif()
    exif[0x0112] = orientation  # Orientation
    exif[0x010F] = "Synthetic Camera"  # Make
    exif[0x8825] = {0x0002: (52.0, 13.0, 0.0)}  # a GPS latitude, which must never leave
    return exif.tobytes()


def photo(size: tuple[int, int] = (3000, 2000), mode: str = "RGB") -> Image.Image:
    """A gradient, so the JPEG has something to encode."""
    image = Image.linear_gradient("L").resize(size)
    if mode == "RGB":
        return Image.merge("RGB", (image, image.transpose(Image.Transpose.FLIP_LEFT_RIGHT), image))
    return image.convert(mode)


def jpeg_with_metadata(size: tuple[int, int] = (3000, 2000), orientation: int = 1) -> bytes:
    """A JPEG carrying Exif (with GPS), an ICC profile and a comment."""
    from PIL import ImageCms

    icc = ImageCms.ImageCmsProfile(ImageCms.createProfile("sRGB")).tobytes()
    out = io.BytesIO()
    photo(size).save(
        out, "JPEG", quality=90, exif=exif_bytes(orientation), icc_profile=icc, comment=b"synthetic"
    )
    return out.getvalue()


def heic_bytes(size: tuple[int, int] = (3000, 2000)) -> bytes:
    """A HEIC with Exif, through pillow-heif in the test environment only."""
    import pillow_heif

    pillow_heif.register_heif_opener()
    out = io.BytesIO()
    photo(size).save(out, "HEIF", quality=80, exif=exif_bytes())
    return out.getvalue()


def jpeg_markers(data: bytes) -> list[int]:
    """The marker bytes of each segment before the scan: 0xE0 is APP0, 0xE1 APP1 (Exif, XMP),
    0xE2 APP2 (ICC), 0xFE a comment."""
    assert data[:2] == b"\xff\xd8", "a JPEG starts with SOI"
    markers, i = [], 2
    while i < len(data):
        assert data[i] == 0xFF
        marker = data[i + 1]
        markers.append(marker)
        if marker == 0xDA:  # start of scan: the entropy-coded data follows
            break
        i += 2 + int.from_bytes(data[i + 2 : i + 4], "big")
    return markers


#: The segments a clean page may carry: JFIF's APP0, quantization and Huffman tables, the frame
#: header, a restart interval and the scan.
CLEAN_MARKERS = {0xE0, 0xDB, 0xC4, 0xC0, 0xC1, 0xC2, 0xDD, 0xDA}


def assert_clean_jpeg(data: bytes) -> Image.Image:
    unexpected = set(jpeg_markers(data)) - CLEAN_MARKERS
    assert not unexpected, f"metadata segments {sorted(hex(m) for m in unexpected)}"
    image = Image.open(io.BytesIO(data))
    assert image.format == "JPEG"
    assert max(image.size) <= 2000
    assert len(data) <= 4 * 1024 * 1024
    assert not image.getexif()
    assert "icc_profile" not in image.info
    return image


def _runs(name: str, *args: str) -> bool:
    path = shutil.which(name)
    if path is None:
        return False
    try:
        return subprocess.run([path, *args], capture_output=True, timeout=60).returncode == 0
    except (OSError, subprocess.TimeoutExpired):
        return False


#: The binary-dependent tests run where the tool is installed and starts: the image always, a dev
#: machine with poppler, LibreOffice or libheif, and CI only for what its runner has.
needs_poppler = pytest.mark.skipif(
    not (_runs("pdftoppm", "-v") and _runs("pdfinfo", "-v")), reason="poppler's pdftoppm and pdfinfo are not installed"
)
needs_soffice = pytest.mark.skipif(not _runs("soffice", "--version"), reason="LibreOffice is not installed")
needs_heif_convert = pytest.mark.skipif(
    not _runs("heif-convert", "--version"), reason="libheif's heif-convert is not installed or does not start"
)
